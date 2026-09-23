"""Runs an AOP workflow_graph. Node kinds: ui_task, data_task, decision, end.
The demo runner stands in for LangGraph; run_ui_task() is the node function LangGraph will wrap."""
import json
import queue
import re
import threading
import time
from datetime import datetime
from pathlib import Path

from .agent import Metrics, run_ui_task
from .browser import Browser

VAR = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def render(obj, variables: dict):
    if isinstance(obj, str):
        whole = VAR.fullmatch(obj.strip())
        if whole:  # a lone {{var}} keeps its type (lists, dicts)
            return _lookup(variables, whole.group(1), obj)
        return VAR.sub(lambda m: str(_lookup(variables, m.group(1), m.group(0))), obj)
    if isinstance(obj, list):
        return [render(x, variables) for x in obj]
    if isinstance(obj, dict):
        return {k: render(v, variables) for k, v in obj.items()}
    return obj


def _lookup(variables, path, default):
    cur = variables
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return default
    return cur


class EventBus:
    """Fan-out of run events to any number of SSE listeners (thread-safe)."""

    def __init__(self):
        self.lock = threading.Lock()
        self.listeners: list[queue.Queue] = []
        self.history: list[dict] = []

    def emit(self, type_: str, data: dict):
        ev = {"type": type_, "t": time.time(), **data}
        brief = {k: v for k, v in data.items() if k not in ("obs", "request", "raw_answers", "metrics")}
        print(f"[{type_}] {json.dumps(brief, default=str)[:300]}", flush=True)
        with self.lock:
            self.history.append(ev)
            for q in self.listeners:
                q.put(ev)

    def subscribe(self) -> queue.Queue:
        q = queue.Queue()
        with self.lock:
            for ev in self.history:
                q.put(ev)
            self.listeners.append(q)
        return q

    def reset(self):
        with self.lock:
            self.history.clear()


class Runtime:
    def __init__(self, cfg, keys, bus: EventBus, frames=None):
        self.cfg, self.keys, self.bus = cfg, keys, bus
        self.metrics = Metrics()
        self.browser = Browser(cfg.get("browser", {}), frames, on_event=self.emit)

    def emit(self, type_, data):
        self.bus.emit(type_, data)


# ---------------- data tasks (deterministic tools) ----------------
def write_excel(data, path, meta: dict) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Captured fields"
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="16181A")
    ws.append(["Field", "Value (as displayed)", "Note"])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for r in data or []:
        ws.append([r.get("field"), r.get("value"), r.get("note")])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 24, 24, 70
    m = wb.create_sheet("Run metadata")
    m.append(["Key", "Value"])
    for c in m[1]:
        c.font, c.fill = head_font, head_fill
    for k, v in meta.items():
        m.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v])
    m.column_dimensions["A"].width, m.column_dimensions["B"].width = 28, 90
    wb.save(path)
    return str(Path(path).resolve())


DATA_TASKS = {"write_excel": write_excel}


def _cond(c: dict, variables: dict) -> bool:
    val = _lookup(variables, c["var"], None)
    op = c.get("op", "eq")
    return {"eq": val == c.get("value"), "ne": val != c.get("value"), "exists": val not in (None, "", [])}[op]


# ---------------- graph execution ----------------
async def run_aop(aop: dict, overrides: dict, rt: Runtime) -> dict:
    schema = aop.get("variables_schema", {})
    variables = {}
    for group in ("required", "optional"):
        for name, spec in schema.get(group, {}).items():
            if "default" in spec:
                variables[name] = spec["default"]
    variables.update({k: v for k, v in (overrides or {}).items() if v not in (None, "")})
    missing = [n for n in schema.get("required", {}) if variables.get(n) in (None, "")]
    if missing:
        raise ValueError(f"missing required variables: {missing}")

    g = aop["workflow_graph"]
    nodes, current = g["nodes"], g["start_node"]
    rt.emit("run_started", {"procedure_id": aop.get("procedure_id"), "version": aop.get("version"),
                            "variables": variables, "nodes": list(nodes)})
    start_url = next((render(n.get("start_url"), variables) for n in nodes.values() if n.get("start_url")), "about:blank")
    await rt.browser.start(start_url)
    result = {"status": "BLOCKED", "reason": "not started"}
    try:
        for _ in range(100):
            raw = nodes[current]
            node = render(raw, variables)
            kind = node["kind"]
            rt.emit("node_started", {"node": current, "kind": kind, "goal": node.get("goal")})
            if kind == "ui_task":
                res = await run_ui_task(node, variables, rt)
                variables.update(res.get("outputs") or {})
                rt.emit("node_finished", {"node": current, **{k: res[k] for k in ("status", "reason")}, "metrics": rt.metrics.snapshot()})
                current = {"DONE": node.get("next"), "HANDOFF": node.get("on_handoff"), "BLOCKED": node.get("on_blocked", node.get("on_handoff"))}[res["status"]]
            elif kind == "data_task":
                fn = DATA_TASKS[node["task"]]
                meta = {"procedure_id": aop.get("procedure_id"), "version": aop.get("version"), "captured_at": datetime.now().isoformat(timespec="seconds"),
                        "url": rt.browser.page.url, **{k: v for k, v in variables.items() if not isinstance(v, (list, dict))},
                        "metrics": rt.metrics.snapshot()}
                path = fn(render(node["data"], variables), node["path"], meta)
                variables[node.get("output_variable", node["task"] + "_path")] = path
                rt.emit("node_finished", {"node": current, "status": "DONE", "reason": f"wrote {path}", "metrics": rt.metrics.snapshot()})
                current = node.get("next")
            elif kind == "decision":
                current = next((r["next"] for r in node.get("rules", []) if _cond(r["when"], variables)), node.get("default_next"))
            elif kind == "end":
                result = {"status": node.get("status", "success"), "reason": node.get("message", "")}
                break
            if not current:
                result = {"status": "error", "reason": "graph has no next node"}
                break
    finally:
        rt.emit("run_finished", {**result, "metrics": rt.metrics.snapshot()})
        await rt.browser.close()
    return result
